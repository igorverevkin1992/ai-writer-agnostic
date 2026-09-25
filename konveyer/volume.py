"""Тома (FR-VL-1…FR-VL-3): сводка тома, закрытие тома, переключение текущего тома.

Рабочая область ведёт один текущий том (манифест `проект.yaml: текущий_том`, `Workspace.volume`):
главы тома 1 — `главы/001`, тома N ≥ 2 — `главы/ТN/001`; выгрузки `выгрузки/` —
всегда текущего тома (`exporter.run_export(volume=…)`); документы канона по тому выбирает
манифест (маркер `Том{N}`/`_Т{N}` в имени).

`close_volume` — закрытие тома (FR-VL-2): все главы «зафиксировано» → рукопись `рукопись/ТомN.md`
(+ `.docx`, если установлен python-docx) и статистика тома (канон не трогают) → снапшот тома в библиотеку
(имя — из каталога типов, через единый конвейер `canonchange.canon_change`) → тег `том-N`.
Переключение текущего тома на следующий — только по подтверждению автора (в CLI).
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from . import accounting, canonchange, exporter, gitops, guard, snapshot
from .config import Config
from .fsm import ChapterState, StatusFileError, all_states
from .paths import Workspace
from .schemas import Act, Brief, Verdict

TAG = "том-{volume}"


def numeric_checks() -> dict[str, str]:
    """{check_id: подпись} метрик Э1 с числовым «actual» — из реестра метрик (§7.5), не из локального списка:
    новая или переименованная в профиле метрика попадает в статистику тома по актам сама."""
    from . import metrics

    out: dict[str, str] = {}
    for m in metrics.REGISTRY.values():
        if m.kind != "метрика" or not m.check_id or not m.unit or m.check_id in out:
            continue
        out[m.check_id] = f"{m.id} ({m.unit})"
    return out


def snapshot_doc_name(volume: int, root: Path | None = None) -> str:
    """Имя документа снапшота тома — из каталога типов (`снапшоты.имя_по_умолчанию`, П-1: в коде имени нет)."""
    from . import catalog

    spec = catalog.load_types(root).get("снапшоты")
    if not (spec and spec.default_name):
        raise LookupError("в каталоге типов нет типа «снапшоты» с «имя_по_умолчанию» — закрытие тома невозможно.")
    return spec.default_name.format(том=int(volume))


def tag_name(volume: int) -> str:
    return TAG.format(volume=int(volume))


# ------------------------------------------------------------------ рукопись и статистика


def prose_files(library: Path, volume: int, root: Path | None = None) -> list[tuple[int, Path]]:
    """Принятые главы тома по документам типа «проза»: [(N, путь)]; макеты не берутся."""
    return exporter.prose_files(library, volume, root)


def _chapter_body(text: str) -> tuple[str, list[str]]:
    """Текст главы без служебной шапки: заголовки `#`, линейки `---` и курсивные пометки в начале файла."""
    paragraphs = [pg.strip() for pg in re.split(r"\n\s*\n", text) if pg.strip()]
    head: list[str] = []
    while paragraphs and (paragraphs[0].startswith("#") or paragraphs[0].startswith(("---", "***"))
                          or re.fullmatch(r"\*[^*]+\*", paragraphs[0])):
        head.append(paragraphs.pop(0))
    return "\n\n".join(paragraphs), head


def _word_count(text: str) -> int:
    return len(re.findall(r"[А-Яа-яЁёA-Za-z0-9]+(?:-[А-Яа-яЁёA-Za-z0-9]+)*", text))


def build_manuscript(ws: Workspace, library: Path, volume: int) -> tuple[Path, Path | None, str | None]:
    """Собирает `рукопись/ТомN.md` (главы по порядку с заголовками) и, если установлен python-docx,
    `рукопись/ТомN.docx`. Возвращает (md, docx | None, подсказка | None)."""
    chapters = prose_files(library, volume, ws.root)
    if not chapters:
        raise FileNotFoundError(
            f"в библиотеке нет принятых глав тома {volume} ({exporter.prose_folder(library, ws.root).name}/"
            f"{exporter.prose_name(library, volume, 1, ws.root)} и далее)."
        )
    briefs = {b.chapter: b for b in _briefs_of(ws, library, volume)}
    lines = [f"# Том {volume}", ""]
    sections: list[tuple[str, list[str]]] = []
    for n, path in chapters:
        body, _ = _chapter_body(path.read_text(encoding="utf-8"))
        title = f"Глава {n}"
        brief = briefs.get(n)
        if brief is not None and brief.date:
            title += f". {brief.date}"
        lines += [f"## {title}", "", body, ""]
        sections.append((title, [pg for pg in re.split(r"\n\s*\n", body) if pg.strip()]))
    md = ws.manuscript / f"Том{volume}.md"
    guard.write_text(md, "\n".join(lines).rstrip("\n") + "\n")
    docx_path = ws.manuscript / f"Том{volume}.docx"
    try:
        import docx  # type: ignore  # python-docx — опциональная зависимость [docx]
    except ImportError:
        return md, None, "python-docx не установлен — собран только .md; `pip install 'konveyer[docx]'` даст и .docx"
    document = docx.Document()
    document.add_heading(f"Том {volume}", level=0)
    for title, paragraphs in sections:
        document.add_heading(title, level=1)
        for pg in paragraphs:
            document.add_paragraph(pg)
    guard.check_write_allowed(docx_path)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(docx_path))
    return md, docx_path, None


def _briefs_of(ws: Workspace, library: Path, volume: int) -> list[Brief]:
    """Поглавник тома: из выгрузок или временного экспорта тома (`accounting.volume_briefs`); нет — пусто."""
    return accounting.volume_briefs(ws, library, volume) or []


def _acts_of(exports: Path) -> list[Act]:
    try:
        return exporter.load_acts(exports)
    except FileNotFoundError:
        return []


def _verdict(ws: Workspace, chapter: int) -> Verdict | None:
    path = ws.chapter_dir(chapter) / "вердикт.json"
    if not path.exists():
        return None
    try:
        return Verdict.model_validate(json.loads(path.read_text(encoding="utf-8")))
    except (ValueError, TypeError):
        return None


def _num(text: str) -> float | None:
    m = re.search(r"-?\d+(?:[.,]\d+)?", text or "")
    return float(m.group().replace(",", ".")) if m else None


@dataclass
class VolumeStats:
    volume: int
    chapters_total: int | None     # глав в поглавнике тома; None — поглавник недоступен
    fixed: list[int]               # «зафиксировано»
    in_work: dict[int, str]        # глава → состояние (не «зафиксировано», не «не-начато»)
    words: dict[int, int]          # слова принятых глав (по документам прозы)
    acts: list[Act]
    act_metrics: dict[int, dict[str, float]]   # акт → {метрика: среднее по главам акта}
    total_metrics: dict[str, float]            # то же по всем главам тома
    cost: float                    # $ по журналы/api.jsonl — все строки тома (сходится с `учёт`)
    calls: int
    author_s: float = 0.0          # время автора по историям глав тома (FR-VL-3, FR-EC-4)
    machine_s: float = 0.0
    missing_docs: list[str] = field(default_factory=list)

    @property
    def words_total(self) -> int:
        return sum(self.words.values())


def volume_stats(ws: Workspace, library: Path, volume: int | None = None) -> VolumeStats:
    """Сводка тома (FR-VL-3): главы по состояниям, слова принятых глав, средние метрики Э1 по актам,
    стоимость и время — из того же учёта, что `konveyer учёт` (`accounting.volume_account`)."""
    volume = ws.volume if volume is None else int(volume)
    vws = ws.for_volume(volume)
    with accounting.volume_exports(ws, library, volume) as exports:
        briefs = accounting.volume_briefs(ws, library, volume) if exports is not None else None
        acts = _acts_of(exports) if exports is not None else []
    fixed: list[int] = []
    in_work: dict[int, str] = {}
    verdicts: dict[int, Verdict] = {}
    for st in all_states(vws):
        if st.state == "зафиксировано":
            fixed.append(st.chapter)
        elif st.state != "не-начато":
            in_work[st.chapter] = st.state
        v = _verdict(vws, st.chapter)
        if v is not None:
            verdicts[st.chapter] = v
    words = {n: _word_count(_chapter_body(p.read_text(encoding="utf-8"))[0]) for n, p in prose_files(library, volume, ws.root)}

    labels = numeric_checks()

    def averages(lo: int, hi: int) -> dict[str, float]:
        sums: dict[str, list[float]] = {}
        for ch, v in verdicts.items():
            if not lo <= ch <= hi:
                continue
            for c in v.checks:
                label = labels.get(c.check_id)
                val = _num(c.actual) if label else None
                if label and val is not None:
                    sums.setdefault(label, []).append(val)
        return {k: round(sum(v) / len(v), 3) for k, v in sorted(sums.items())}

    act_metrics = {a.act: m for a in acts if (m := averages(a.from_chapter, a.to_chapter))}
    acc = accounting.volume_account(ws, volume, chapters_total=len(briefs) if briefs is not None else 0)
    return VolumeStats(
        volume=volume, chapters_total=len(briefs) if briefs is not None else None, fixed=sorted(fixed),
        in_work=dict(sorted(in_work.items())), words=words, acts=acts, act_metrics=act_metrics,
        total_metrics=averages(0, 10**6), cost=round(acc.cost, 2), calls=acc.calls,
        author_s=acc.author_s, machine_s=acc.machine_s, missing_docs=exporter.missing_volume_docs(library, volume, ws.root),
    )


def render_stats(stats: VolumeStats) -> str:
    from . import timing

    lines = [f"# Том {stats.volume} — статистика", ""]
    plan = stats.chapters_total if stats.chapters_total is not None else "поглавник недоступен"
    lines.append(f"- Глав в поглавнике: {plan}; зафиксировано: {len(stats.fixed)}; в работе: {len(stats.in_work)}")
    lines.append(f"- Слов в принятых главах: {stats.words_total}")
    lines.append(f"- Вызовов моделей: {stats.calls}; оценка стоимости: ${stats.cost:.2f}")
    lines.append(f"- Время автора: {timing.fmt_minutes(stats.author_s)}; машинное: {timing.fmt_minutes(stats.machine_s)}")
    if stats.in_work:
        lines += ["", "## В работе", ""]
        lines += [f"- Глава {n}: {s}" for n, s in stats.in_work.items()]
    if stats.words:
        lines += ["", "## Объём глав", "", "| Глава | Слов |", "|---|---|"]
        lines += [f"| {n} | {w} |" for n, w in sorted(stats.words.items())]
    if stats.total_metrics:
        lines += ["", "## Метрики Э1 (средние по вердиктам глав)", ""]
        lines.append("- Весь том: " + "; ".join(f"{k}: {v:g}" for k, v in stats.total_metrics.items()))
        for a in stats.acts:
            m = stats.act_metrics.get(a.act, {})
            words = sum(w for n, w in stats.words.items() if a.from_chapter <= n <= a.to_chapter)
            desc = "; ".join(f"{k}: {v:g}" for k, v in m.items()) or "вердиктов нет"
            lines.append(f"- Акт {a.act} «{a.title}» (гл. {a.from_chapter}–{a.to_chapter}, {words} слов): {desc}")
    return "\n".join(lines) + "\n"


def write_stats(ws: Workspace, stats: VolumeStats) -> Path:
    path = ws.manuscript / f"Том{stats.volume}_статистика.md"
    guard.write_text(path, render_stats(stats))
    return path


# ------------------------------------------------------------------ закрытие и переключение


@dataclass
class CloseResult:
    volume: int
    snapshot_doc: Path
    commit: str | None
    tag: str | None
    manuscript_md: Path
    manuscript_docx: Path | None
    docx_hint: str | None
    stats_path: Path
    messages: list[str] = field(default_factory=list)


def unfixed_chapters(ws: Workspace, library: Path, volume: int) -> list[str]:
    """Главы тома, которые ещё не «зафиксировано» (по поглавнику тома и папкам глав): «гл. N (состояние)»."""
    vws = ws.for_volume(volume)
    states: dict[int, str] = {}
    for n, _ in vws.chapter_dirs():
        try:
            states[n] = ChapterState(vws, n).state
        except StatusFileError:
            states[n] = "повреждено"
    for b in _briefs_of(ws, library, volume):
        states.setdefault(b.chapter, "не-начато")
    return [f"гл. {n} ({s})" for n, s in sorted(states.items()) if s != "зафиксировано"]


def missing_prose(ws: Workspace, library: Path, volume: int) -> list[str]:
    """Зафиксированные главы тома, у которых в библиотеке нет документа прозы (удалён, переименован,
    другой шаблон имени в типе «проза»): рукопись без них молча была бы неполной (FR-VL-2)."""
    have = {n for n, _ in prose_files(library, volume, ws.root)}
    vws = ws.for_volume(volume)
    out = []
    for st in all_states(vws):
        if st.state == "зафиксировано" and st.chapter not in have:
            out.append(f"гл. {st.chapter} ({exporter.prose_name(library, volume, st.chapter, ws.root)})")
    return out


def close_volume(ws: Workspace, cfg: Config, library: Path, volume: int, *, again: bool = False,
                 author_confirmed: bool) -> CloseResult:
    """Закрывает том (FR-VL-2). Все проверки и всё, что не трогает канон (рукопись, статистика), — ДО записи
    в библиотеку; затем снапшот тома через `canonchange.canon_change` (коммит) и тег `том-N`: сбой на рукописи
    не оставляет том «полузакрытым». Как и приёмка, требует библиотеку под git (иначе нет ни коммита, ни тега).
    Выгрузки должны быть тома `volume` (закрывается текущий том; другой — сначала `konveyer том открыть N`)."""
    if not author_confirmed:
        raise PermissionError("закрытие тома без подтверждения автора запрещено (FR-K2, Д-8).")
    if volume != ws.volume:
        raise RuntimeError(
            f"закрывается только текущий том рабочей области (сейчас том {ws.volume}); "
            f"переключитесь: `konveyer том открыть {volume}`."
        )
    if not gitops.is_repo(library):
        raise RuntimeError(
            "библиотека не под git — закрытие тома невозможно (снапшот тома коммитится и помечается тегом, "
            "как приёмка главы). Инициализируйте репозиторий в библиотеке (git init; git add -A; git commit), затем повторите."
        )
    exporter.run_export(library, ws.exports, ws.logs, volume, ws.root)
    pending = unfixed_chapters(ws, library, volume)
    if pending:
        raise RuntimeError(
            f"том {volume} нельзя закрыть: не зафиксированы {', '.join(pending)}. "
            "Доведите главы до «зафиксировано» (`konveyer канон N --применить`)."
        )
    if not _briefs_of(ws, library, volume):
        raise RuntimeError(f"в поглавнике нет глав тома {volume} — закрывать нечего.")
    lost = missing_prose(ws, library, volume)
    if lost:
        raise RuntimeError(
            f"том {volume} нельзя закрыть: у зафиксированных глав нет документа прозы в библиотеке — {', '.join(lost)}. "
            "Верните файлы прозы (или откатите главы), затем повторите."
        )
    doc = library / snapshot_doc_name(volume, ws.root)
    messages: list[str] = []
    if doc.exists() and not again:
        raise RuntimeError(
            f"снапшот {doc.name} уже есть в библиотеке; чтобы переписать его — `konveyer том закрыть {volume} --заново`."
        )
    # 1. рукопись и статистика — канон не трогают; их сбой (нет папки, нет прав, python-docx) ничего не портит
    draft = snapshot.build_snapshot(ws, volume)
    md, docx_path, hint = build_manuscript(ws, library, volume)
    stats_path = write_stats(ws, volume_stats(ws, library, volume))
    # 2. снапшот — в канон через единый конвейер (FR-K3), затем тег
    text = draft.read_text(encoding="utf-8").replace(
        "Черновик сгенерирован конвейером; вносится в канон правкой библиотеки + `konveyer канон-коммит`.",
        f"Внесён в канон командой `konveyer том закрыть {volume}` (срез мира на конец тома).",
    )
    result = canonchange.canon_change(
        ws, cfg, library, lambda: guard.write_text(doc, text),
        f"[том {volume}] снапшот: {doc.name} (закрытие тома)",
        commit=True, author_confirmed=True, action=f"закрытие тома {volume}",
    )
    messages.append(result.message)
    tag: str | None = None
    sha = result.commit or gitops.head(library)
    name = tag_name(volume)
    try:
        existing = gitops.tags(library, name)
        if existing and not again:
            messages.append(f"тег {name} уже стоит — оставлен как есть.")
            tag = name
        else:
            if existing:
                gitops.delete_tag(library, name)
            tag = gitops.tag(library, name, sha)
    except (RuntimeError, FileNotFoundError) as e:
        messages.append(f"тег {name} не поставлен: {e}")
    return CloseResult(
        volume=volume, snapshot_doc=doc, commit=result.commit, tag=tag, manuscript_md=md,
        manuscript_docx=docx_path, docx_hint=hint, stats_path=stats_path, messages=messages,
    )


def open_volume(ws: Workspace, library: Path, volume: int) -> list[str]:
    """Проверяет, что обязательные документы тома есть в библиотеке; возвращает список недостающих (пусто — можно
    переключать `config.volume`, это делает CLI через `config.set_volume`). Документы модулей — см. `volume_warnings`."""
    if int(volume) < 1:
        raise ValueError(f"номер тома должен быть ≥ 1, получено: {volume}.")
    return exporter.missing_volume_docs(library, int(volume), ws.root, only_required=True)


def volume_warnings(ws: Workspace, library: Path, volume: int) -> list[str]:
    """Потомные документы включённых модулей, которых у тома нет (предупреждение, не отказ)."""
    required = set(open_volume(ws, library, volume))
    return [m for m in exporter.missing_volume_docs(library, int(volume), ws.root) if m not in required]
